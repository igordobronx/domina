"""Variantes de prompt para o teste de calibracao (nao fazem parte do projeto)."""
import random
from typing import Literal

from pathlib import Path

import openpyxl
from pydantic import BaseModel, Field

from app.correcao import PROMPT as V2, ResultadoIA

PLANILHA = Path(__file__).resolve().parents[2] / "essays" / "base_redacoes_primeira.xlsm"


def amostra(n1000: int, nuol: int, seed: int = 42):
    # mesma amostra (seed 42) usada por calibra.py, extraida pra ca para que
    # qualquer backend novo (ex.: Groq) teste exatamente as mesmas redacoes.
    rows = [r for r in list(openpyxl.load_workbook(PLANILHA, read_only=True, data_only=True)["Redações"].iter_rows(values_only=True))[1:] if r[0]]
    mil = [r for r in rows if str(r[10]).startswith(("Cartilha", "Toda"))]
    uol = [r for r in rows if str(r[10]).startswith("Banco") and r[3] and len(str(r[3]).split()) >= 120]
    rnd = random.Random(seed)
    fracas = [r for r in uol if float(r[9]) <= 350]
    medias = [r for r in uol if 400 <= float(r[9]) <= 600]
    boas = [r for r in uol if float(r[9]) >= 700]
    k = nuol // 3
    return [("mil", r) for r in rnd.sample(mil, n1000)] + [("uol", r) for r in rnd.sample(fracas, k) + rnd.sample(medias, k) + rnd.sample(boas, nuol - 2 * k)]

# --- v3a: pensa antes de dar a nota, conta desvios na C1, usa a escala inteira -------------
EXTRA_V3A = """
Antes de dar cada nota, escreva a análise (campo "analise") e só depois escolha o nível.

Para a C1, liste os desvios claros encontrados (concordância, regência, crase, pontuação que
compromete o sentido, ortografia, registro informal) e use a quantidade como referência,
lembrando que a banca do INEP tolera desvios pontuais em textos bem escritos:
  200: até 2 desvios, sem reincidência.   160: 3 a 5 desvios.   120: 6 a 9 desvios ou reincidência.
   80: 10 ou mais desvios, ou desvios que atrapalham a leitura.   40/0: domínio precário.

Use a escala inteira. Não puxe as notas para o meio:
- texto bem estruturado, com repertório pertinente e proposta completa deve receber 160 ou 200;
- texto curto, superficial, sem repertório, com cópia da proposta ou sem proposta de
  intervenção deve receber 80 ou menos nas competências afetadas, mesmo se a gramática for boa.
"""
V3A = V2.replace("Tema proposto: {tema}", EXTRA_V3A + "\nTema proposto: {tema}")


class CompetenciaV3(BaseModel):
    codigo: Literal["C1", "C2", "C3", "C4", "C5"]
    analise: str  # vem antes da nota: o modelo raciocina e depois decide
    nota: int


class ResultadoV3(BaseModel):
    competencias: list[CompetenciaV3] = Field(min_length=5, max_length=5)
    feedback_geral: str
    confianca: float = Field(ge=0, le=1)


def para_v2(r: ResultadoV3) -> ResultadoIA:
    # reaproveita as validacoes da v2 (niveis do ENEM, C1..C5 uma vez cada)
    return ResultadoIA.model_validate({
        "competencias": [{"codigo": c.codigo, "analise": c.analise, "nota": c.nota, "justificativa": c.analise} for c in r.competencias],
        "feedback_geral": r.feedback_geral, "confianca": r.confianca})


# --- v3b: v3a + 3 redacoes de referencia (fora da amostra de teste) ------------------------
# notas do UOL (passo 50) arredondadas para o nivel do INEP mais proximo (empate para baixo)
ANCORAS = {"0008": [80, 40, 40, 40, 0], "0021": [160, 80, 80, 80, 80], "0023": [160, 200, 200, 160, 160]}


def _v3b():
    rows = {str(r[0]): r for r in list(openpyxl.load_workbook(PLANILHA, read_only=True,
            data_only=True)["Redações"].iter_rows(values_only=True))[1:] if r[0]}
    blocos = []
    for i, (rid, notas) in enumerate(ANCORAS.items(), 1):
        r = rows[rid]
        blocos.append(f"Exemplo {i} (tema: {r[2]}) - notas da banca: " +
                      ", ".join(f"C{k+1}={n}" for k, n in enumerate(notas)) + f" (total {sum(notas)})\n\"\"\"\n{r[3]}\n\"\"\"")
    exemplos = ("\nRedações de referência já corrigidas, para calibrar a escala (NÃO são a redação a corrigir):\n\n"
                + "\n\n".join(blocos) + "\n")
    return V3A.replace("Tema proposto: {tema}", exemplos.replace("{", "{{").replace("}", "}}") + "\nTema proposto: {tema}")


VARIANTES = {
    "v2": (V2, ResultadoIA, lambda r: r),
    "v3a": (V3A, ResultadoV3, para_v2),
    "v3b": (None, ResultadoV3, para_v2),  # montado sob demanda (le a planilha)
}


def pegar(nome):
    prompt, schema, conv = VARIANTES[nome]
    return (_v3b() if nome == "v3b" else prompt), schema, conv
